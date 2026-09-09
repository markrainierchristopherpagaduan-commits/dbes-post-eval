"""
report.py — generates a Word (.docx) post-evaluation report for an
activity, combining the quantitative rating summary, per-speaker
breakdown, and AI-generated qualitative theme summaries.
"""

from __future__ import annotations

import io

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

import db

ACCENT = RGBColor(0x7A, 0x1F, 0x2B)


def _add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = ACCENT
    return h


def _rating_table(doc, rows, headers):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val) if val is not None else "—"
    return table


def build_report(activity_id: str) -> bytes:
    activity = db.get_activity(activity_id)
    if not activity:
        raise ValueError("Activity not found")

    speakers = db.get_activity_speakers(activity_id)
    rating_summary = db.get_rating_summary(activity_id)
    speaker_averages = db.get_speaker_rating_averages(activity_id)
    ai_summaries = db.get_ai_summaries(activity_id)
    n_responses = db.count_responses(activity_id)

    doc = Document()

    # --- Title page ---------------------------------------------------
    title = doc.add_heading("Post-Evaluation Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = ACCENT

    subtitle = doc.add_paragraph(activity["title"])
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(16)
    subtitle.runs[0].bold = True

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_bits = [b for b in [activity["activity_type"], activity["activity_date"], activity["venue"],
                              activity["hosting_school"] or activity["participants"]] if b]
    meta.add_run(" · ".join(meta_bits)).italic = True

    doc.add_paragraph()

    # --- Overview -------------------------------------------------------
    _add_heading(doc, "Overview", level=1)
    overview = doc.add_paragraph()
    overview.add_run(f"Total responses collected: ").bold = True
    overview.add_run(str(n_responses))
    if speakers:
        sp_line = doc.add_paragraph()
        sp_line.add_run("Speaker(s): ").bold = True
        sp_line.add_run(", ".join(
            sp["name"] + (f' ({sp["topic"]})' if sp["topic"] else "") for sp in speakers
        ))

    # --- Quantitative summary -------------------------------------------
    _add_heading(doc, "Quantitative Summary", level=1)
    if rating_summary:
        rows = [
            (r["category"], r["question_text"], f'{r["avg_rating"]:.2f}' if r["avg_rating"] is not None else "—", r["n"])
            for r in rating_summary
        ]
        _rating_table(doc, rows, ["Category", "Question", "Avg. Rating (1–5)", "# Responses"])
    else:
        doc.add_paragraph("No rating-type questions were configured for this activity.")

    # --- Per-speaker breakdown -------------------------------------------
    if speakers:
        _add_heading(doc, "Per-Speaker Ratings", level=1)
        rows = [
            (sa["name"], sa["topic"] or "—", f'{sa["avg_rating"]:.2f}' if sa["avg_rating"] is not None else "—", sa["n"])
            for sa in speaker_averages
        ]
        _rating_table(doc, rows, ["Speaker", "Topic", "Overall Avg. Rating (1–5)", "# Ratings"])

    # --- Qualitative summary ---------------------------------------------
    _add_heading(doc, "Qualitative Summary", level=1)
    if ai_summaries:
        for s in ai_summaries:
            doc.add_heading(s["category"], level=2)
            doc.add_paragraph(s["summary_text"])
    else:
        doc.add_paragraph(
            "No qualitative summary has been written yet for this activity. "
            "Add one from the Activity Results page before exporting a final report."
        )

    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.add_run(
        "Generated automatically by the DBES Post-Evaluation System."
    ).italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
